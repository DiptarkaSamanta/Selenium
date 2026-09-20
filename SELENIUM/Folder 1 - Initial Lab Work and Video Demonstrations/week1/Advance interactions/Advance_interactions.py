# =========================================================================
# IMPLEMENTATION: ADVANCE INTERACTIONS & DATA HANDLING
# Website: https://practice.softwaretestingmentor.com/
# Syllabus Topics Covered:
# 1. Different API to read data from excel, Json
# 2. Read data from properties, CSV, XML file
# 3. Reading from excel sheet / Json
# 4. Writing Data into excel sheet / Json
# 5. Mouse Hover Actions
# 6. Executing JavaScript Commands
# =========================================================================

import os
import json
import csv
import xml.etree.ElementTree as ET
import configparser
import time
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.common.action_chains import ActionChains

# Import openpyxl for Excel handling
try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

# =========================================================================
# 1 - 4. DATA READING & WRITING (EXCEL, JSON, PROPERTIES, CSV, XML)
# =========================================================================
def demo_data_files():
    print("==========================================================================")
    print("1 - 4. DATA READING & WRITING OPERATIONS")
    print("==========================================================================")

    current_dir = os.path.dirname(os.path.abspath(__file__))

    # -------------------------------------------------------------------------
    # 1 & 3: READING FROM JSON
    # -------------------------------------------------------------------------
    json_path = os.path.join(current_dir, "sample_config.json")
    json_data = {
        "website": "https://practice.softwaretestingmentor.com/",
        "credentials": {"username": "DemoUser", "password": "DemoPassword123"}
    }
    with open(json_path, "w") as f:
        json.dump(json_data, f, indent=4)

    with open(json_path, "r") as f:
        read_json = json.load(f)
    print("-> [JSON READ] Target Website:", read_json["website"])
    print("-> [JSON READ] Username:", read_json["credentials"]["username"])

    # -------------------------------------------------------------------------
    # 4: WRITING DATA INTO JSON
    # -------------------------------------------------------------------------
    json_out_path = os.path.join(current_dir, "test_output.json")
    output_data = {"test_status": "PASSED", "execution_time_sec": 4.5}
    with open(json_out_path, "w") as f:
        json.dump(output_data, f, indent=4)
    print("-> [JSON WRITE] Wrote test status output to:", json_out_path)

    # -------------------------------------------------------------------------
    # 2: READING FROM PROPERTIES (CONFIG.INI) FILE
    # -------------------------------------------------------------------------
    prop_path = os.path.join(current_dir, "config.properties")
    config = configparser.ConfigParser()
    config["DEFAULT"] = {"browser": "chrome", "timeout": "10"}
    config["TEST"] = {"url": "https://practice.softwaretestingmentor.com/login"}
    with open(prop_path, "w") as f:
        config.write(f)

    read_cfg = configparser.ConfigParser()
    read_cfg.read(prop_path)
    print("-> [PROPERTIES READ] Browser:", read_cfg["DEFAULT"]["browser"])
    print("-> [PROPERTIES READ] URL:", read_cfg["TEST"]["url"])

    # -------------------------------------------------------------------------
    # 2: READING FROM CSV FILE
    # -------------------------------------------------------------------------
    csv_path = os.path.join(current_dir, "test_data.csv")
    with open(csv_path, "w", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["Username", "Password"])
        writer.writerow(["user1", "pass1"])
        writer.writerow(["user2", "pass2"])

    print("-> [CSV READ] Reading CSV rows:")
    with open(csv_path, "r") as f:
        reader = csv.reader(f)
        for row in reader:
            print("   ", row)

    # -------------------------------------------------------------------------
    # 2: READING FROM XML FILE
    # -------------------------------------------------------------------------
    xml_path = os.path.join(current_dir, "config.xml")
    xml_content = "<config><environment name='QA'><url>https://practice.softwaretestingmentor.com/</url></environment></config>"
    with open(xml_path, "w") as f:
        f.write(xml_content)

    tree = ET.parse(xml_path)
    root = tree.getroot()
    env_name = root.find("environment").attrib["name"]
    env_url = root.find("environment/url").text
    print(f"-> [XML READ] Environment: {env_name} | URL: {env_url}")

    # -------------------------------------------------------------------------
    # 1, 3 & 4: EXCEL READING & WRITING (OPENPYXL)
    # -------------------------------------------------------------------------
    if HAS_OPENPYXL:
        excel_path = os.path.join(current_dir, "test_excel_data.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "LoginData"
        ws.append(["Username", "Password", "Status"])
        ws.append(["DemoUser", "DemoPassword123", "PASSED"])
        wb.save(excel_path)
        print("-> [EXCEL WRITE] Saved Excel workbook to:", excel_path)

        wb_read = openpyxl.load_workbook(excel_path)
        sheet = wb_read["LoginData"]
        print("-> [EXCEL READ] Reading Excel rows:")
        for r in range(1, sheet.max_row + 1):
            vals = [sheet.cell(row=r, column=c).value for c in range(1, sheet.max_column + 1)]
            print("   ", vals)
    else:
        print("-> openpyxl package not installed. Skipping Excel operation.")


# =========================================================================
# 5 & 6. MOUSE HOVER ACTIONS & JAVASCRIPT EXECUTOR
# =========================================================================
def demo_selenium_actions():
    print("\n==========================================================================")
    print("5 & 6. MOUSE HOVER ACTIONS & JAVASCRIPT EXECUTOR")
    print("Website: https://practice.softwaretestingmentor.com/")
    print("==========================================================================")

    options = webdriver.ChromeOptions()
    options.add_argument("--disable-blink-features=AutomationControlled")

    driver = webdriver.Chrome(options=options)
    driver.maximize_window()

    def open_url(url):
        for _ in range(3):
            try:
                driver.get(url)
                time.sleep(3)
                return
            except Exception:
                time.sleep(2)

    try:
        open_url("https://practice.softwaretestingmentor.com/login")

        # -------------------------------------------------------------------------
        # 5. MOUSE HOVER ACTIONS USING ActionChains
        # -------------------------------------------------------------------------
        print("\n--- 5. Mouse Hover Actions ---")
        login_btn = driver.find_element(By.ID, "login-btn")

        actions = ActionChains(driver)
        actions.move_to_element(login_btn).perform()
        print("Mouse hover performed on Sign In button successfully!")
        time.sleep(1)

        # -------------------------------------------------------------------------
        # 6. EXECUTING JAVASCRIPT COMMANDS
        # -------------------------------------------------------------------------
        print("\n--- 6. Executing JavaScript Commands ---")

        # A. Fetch Document Title via JS
        js_title = driver.execute_script("return document.title;")
        print("JS Document Title:", js_title)

        # B. Scroll Page via JS
        driver.execute_script("window.scrollBy(0, 200);")
        print("JS Scrolled down by 200px.")
        time.sleep(1)

        # C. Set Input Field Value via JS
        user_field = driver.find_element(By.ID, "username")
        driver.execute_script("arguments[0].value='JS_Input_Value';", user_field)
        print("JS Populated Username field:", user_field.get_attribute("value"))
        time.sleep(1)

        # D. Highlight Element via JS (Add Red Border)
        driver.execute_script("arguments[0].style.border='3px solid red';", user_field)
        print("JS Highlighted Username field with red border.")
        time.sleep(1)

        # E. Click Element via JS
        driver.execute_script("arguments[0].click();", login_btn)
        print("JS Executed Click on Sign In button.")
        time.sleep(1)

    finally:
        time.sleep(2)
        driver.quit()
        print("\nBrowser closed successfully.")


if __name__ == "__main__":
    demo_data_files()
    demo_selenium_actions()
